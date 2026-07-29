"""File and SQL loaders that feed the TabularView spine.

All loaders return a :class:`~vizly.data.TabularView` (typically
:class:`~vizly.data.ColumnarTable`). Chart factories accept the result
directly — callers do not need to build a pandas DataFrame.

SQL uses SQLAlchemy (base dependency). Database drivers and external
dialects remain the caller's install. vizly executes the SQL you provide
through your Engine/Connection/URL; it does not open live vendor APIs
beyond that.
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence, TextIO, Union

from vizly.data import ColumnarTable, DataError, DataLike, TabularView, as_tabular

PathLike = Union[str, Path]
SourceLike = Union[PathLike, TextIO, bytes]


def from_records(rows: Sequence[Mapping[str, Any]]) -> TabularView:
    """Build a tabular view from a sequence of row mappings."""
    return ColumnarTable.from_records(list(rows))


def from_columnar(data: Mapping[str, Sequence[Any]]) -> TabularView:
    """Build a tabular view from a columnar mapping."""
    return ColumnarTable(data)


def _read_text(source: SourceLike, *, encoding: str = "utf-8") -> str:
    if isinstance(source, (str, Path)):
        path = Path(source)
        return path.read_text(encoding=encoding)
    if isinstance(source, bytes):
        return source.decode(encoding)
    return source.read()


def from_csv(
    source: SourceLike,
    *,
    encoding: str = "utf-8",
    delimiter: str = ",",
) -> TabularView:
    """Load CSV from a path, text stream, or bytes."""
    text = _read_text(source, encoding=encoding)
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    rows = [dict(row) for row in reader]
    # Attempt numeric coercion for common chart payloads
    return _coerce_numeric_strings(ColumnarTable.from_records(rows))


def from_tsv(source: SourceLike, *, encoding: str = "utf-8") -> TabularView:
    """Load TSV (tab-delimited) from a path, text stream, or bytes."""
    return from_csv(source, encoding=encoding, delimiter="\t")


def from_json(
    source: Union[SourceLike, Any],
    *,
    encoding: str = "utf-8",
    orient: str = "records",
) -> TabularView:
    """Load JSON as records, columnar object, or nested list-of-objects.

    ``orient``:
    - ``records`` (default): ``[{...}, ...]`` or auto-detect
    - ``columnar``: ``{"col": [...], ...}``
    - ``auto``: detect records vs columnar
    """
    if isinstance(source, (str, Path, bytes)) or hasattr(source, "read"):
        if isinstance(source, str) and not Path(source).exists() and source.lstrip().startswith(("[", "{")):
            payload = json.loads(source)
        elif isinstance(source, bytes):
            payload = json.loads(source.decode(encoding))
        elif isinstance(source, (str, Path)):
            path = Path(source)
            if path.exists():
                payload = json.loads(path.read_text(encoding=encoding))
            else:
                payload = json.loads(str(source))
        else:
            payload = json.loads(source.read())  # type: ignore[union-attr]
    else:
        payload = source

    if orient == "columnar":
        if not isinstance(payload, Mapping):
            raise DataError("from_json(orient='columnar') expects a JSON object.")
        return ColumnarTable(payload)  # type: ignore[arg-type]

    if orient == "records":
        if isinstance(payload, list):
            return ColumnarTable.from_records(payload)
        if isinstance(payload, Mapping):
            # Nested: {"data": [...]} common pattern
            for key in ("data", "records", "rows", "items", "results"):
                if key in payload and isinstance(payload[key], list):
                    return ColumnarTable.from_records(payload[key])
            # Columnar object mistaken for records — accept as columnar
            return ColumnarTable(payload)  # type: ignore[arg-type]
        raise DataError(
            "from_json(orient='records') expects a JSON array of objects "
            "or an object with a data/records/rows list."
        )

    # auto
    if isinstance(payload, list):
        return ColumnarTable.from_records(payload)
    if isinstance(payload, Mapping):
        vals = list(payload.values())
        if vals and all(isinstance(v, list) for v in vals):
            return ColumnarTable(payload)  # type: ignore[arg-type]
        for key in ("data", "records", "rows", "items", "results"):
            if key in payload and isinstance(payload[key], list):
                return ColumnarTable.from_records(payload[key])
        return ColumnarTable.from_records([dict(payload)])
    raise DataError(f"Unsupported JSON payload type {type(payload).__name__}.")


def from_excel(
    source: PathLike,
    *,
    sheet_name: Union[str, int, None] = 0,
) -> TabularView:
    """Load an Excel workbook sheet via openpyxl (base dependency)."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "from_excel requires openpyxl (a vizly base dependency)."
        ) from exc

    path = Path(source)
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        if sheet_name is None:
            ws = wb.active
        elif isinstance(sheet_name, int):
            ws = wb.worksheets[sheet_name]
        else:
            ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return ColumnarTable({})
        keys = [str(h) if h is not None else f"col{i}" for i, h in enumerate(header)]
        cols: dict = {k: [] for k in keys}
        for row in rows_iter:
            if row is None or all(c is None for c in row):
                continue
            for i, key in enumerate(keys):
                cols[key].append(row[i] if i < len(row) else None)
        return ColumnarTable(cols)
    finally:
        wb.close()


def from_sql(
    statement: str,
    bind: Any = None,
    *,
    url: Optional[str] = None,
    params: Optional[Mapping[str, Any]] = None,
) -> TabularView:
    """Execute SQL via SQLAlchemy and return a TabularView.

    Pass either ``bind=`` (Engine or Connection) or ``url=`` (SQLAlchemy URL
    string). DBAPI drivers and external dialects are the caller's
    responsibility — if SQLAlchemy can connect, vizly can chart the result.

    Example::

        from sqlalchemy import create_engine
        import vizly as vz

        engine = create_engine("sqlite:///:memory:")
        table = vz.from_sql("SELECT region, sales FROM t", bind=engine)
        vz.bar(table, x="region", y="sales")
    """
    try:
        from sqlalchemy import create_engine, text
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "from_sql requires sqlalchemy (a vizly base dependency)."
        ) from exc

    if bind is None and url is None:
        raise DataError("from_sql requires bind= (Engine/Connection) or url=.")
    if bind is not None and url is not None:
        raise DataError("from_sql accepts bind= or url=, not both.")

    engine_or_conn = bind if bind is not None else create_engine(url)
    owns_engine = bind is None

    try:
        # Connection.execute vs Engine.connect
        if hasattr(engine_or_conn, "execute") and not hasattr(engine_or_conn, "connect"):
            result = engine_or_conn.execute(text(statement), params or {})
            keys = list(result.keys())
            cols: dict = {k: [] for k in keys}
            for row in result.mappings():
                for k in keys:
                    cols[k].append(row[k])
            return ColumnarTable(cols)
        with engine_or_conn.connect() as conn:
            result = conn.execute(text(statement), params or {})
            keys = list(result.keys())
            cols = {k: [] for k in keys}
            for row in result.mappings():
                for k in keys:
                    cols[k].append(row[k])
            return ColumnarTable(cols)
    finally:
        if owns_engine and hasattr(engine_or_conn, "dispose"):
            engine_or_conn.dispose()


def _coerce_numeric_strings(table: ColumnarTable) -> ColumnarTable:
    """Best-effort: coerce all-numeric string columns to float/int."""
    columnar = table.to_columnar()
    out: dict = {}
    for key, values in columnar.items():
        non_null = [v for v in values if v is not None and v != ""]
        if not non_null:
            out[key] = values
            continue
        if not all(isinstance(v, str) for v in non_null):
            out[key] = values
            continue
        coerced = []
        ok = True
        for v in values:
            if v is None or v == "":
                coerced.append(None)
                continue
            try:
                if isinstance(v, str) and "." not in v and "e" not in v.lower():
                    coerced.append(int(v))
                else:
                    coerced.append(float(v))
            except (TypeError, ValueError):
                ok = False
                break
        out[key] = coerced if ok else values
    return ColumnarTable(out)


def as_chart_data(data: DataLike) -> TabularView:
    """Normalize any supported chart input to TabularView (alias of as_tabular)."""
    return as_tabular(data)
